"""Verificar que las celdas F3, E4 e I3 tienen los valores correctos."""
import openpyxl
from pathlib import Path

def verificar_celdas_adicionales(filepath):
    """Verifica que las celdas especiales tienen los valores esperados."""
    if not filepath.exists():
        print(f"❌ Archivo no encontrado: {filepath}")
        return False
    
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active
    
    print("=" * 70)
    print(f"VERIFICACIÓN DE CELDAS ADICIONALES")
    print(f"Archivo: {filepath}")
    print("=" * 70)
    
    # Verificar celda F3 (nombre trabajador)
    f3_value = ws.cell(row=3, column=6).value  # Columna F
    print(f"\n📍 Celda F3 (Nombre trabajador):")
    print(f"   Valor: {f3_value}")
    print(f"   Estado: {'✅ OK' if f3_value else '❌ Vacía'}")
    
    # Verificar celda E4 (rango de fechas)
    e4_value = ws.cell(row=4, column=5).value  # Columna E
    print(f"\n📍 Celda E4 (Rango de fechas):")
    print(f"   Valor: {e4_value}")
    print(f"   Estado: {'✅ OK' if e4_value and '-' in str(e4_value) else '❌ Vacía o formato incorrecto'}")
    
    # Verificar celda I3 (payment type con periodo)
    i3_value = ws.cell(row=3, column=9).value  # Columna I
    print(f"\n📍 Celda I3 (Payment type con periodo):")
    print(f"   Valor: {i3_value}")
    print(f"   Estado: {'✅ OK' if i3_value and 'Corporate Card' in str(i3_value) else '❌ Vacía o formato incorrecto'}")
    
    # Verificar también los datos en las filas 8-10
    print(f"\n📍 Datos en filas 8-10:")
    for row in range(8, 11):
        id_val = ws.cell(row=row, column=3).value
        fecha_val = ws.cell(row=row, column=4).value
        amount_val = ws.cell(row=row, column=9).value
        print(f"   Fila {row}: ID={id_val}, Fecha={fecha_val}, Amount={amount_val}")
    
    # Verificar fórmula de suma
    formula_val = ws.cell(row=11, column=9).value
    print(f"\n📍 Fórmula de suma (fila 11, columna I):")
    print(f"   Valor: {formula_val}")
    print(f"   Estado: {'✅ OK' if formula_val and str(formula_val).startswith('=SUM') else '❌ No encontrada'}")
    
    wb.close()
    
    print("\n" + "=" * 70)
    
    # Resumen
    all_ok = (
        f3_value is not None and
        e4_value is not None and '-' in str(e4_value) and
        i3_value is not None and 'Corporate Card' in str(i3_value) and
        formula_val is not None and str(formula_val).startswith('=SUM')
    )
    
    if all_ok:
        print("✅ TODAS LAS VERIFICACIONES PASARON")
    else:
        print("⚠️  ALGUNAS VERIFICACIONES FALLARON")
    
    print("=" * 70)
    
    return all_ok

if __name__ == "__main__":
    archivo = Path("doc/examples/test_export_resultado.xlsx")
    verificar_celdas_adicionales(archivo)

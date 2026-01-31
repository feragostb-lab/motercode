"""Debug script para verificar la celda F3."""
import openpyxl
from pathlib import Path

# Verificar el archivo generado
archivo = Path("doc/examples/test_export_resultado.xlsx")

if archivo.exists():
    wb = openpyxl.load_workbook(archivo)
    ws = wb.active
    
    print("=" * 70)
    print("VERIFICACIÓN DETALLADA DE CELDAS")
    print("=" * 70)
    
    # Verificar F3
    f3_value = ws.cell(row=3, column=6).value
    print(f"\n🔍 Celda F3 (fila 3, columna 6):")
    print(f"   Valor: '{f3_value}'")
    print(f"   Tipo: {type(f3_value)}")
    print(f"   Es None: {f3_value is None}")
    print(f"   Es vacío: {f3_value == ''}")
    
    # Verificar todas las celdas de la fila 3
    print(f"\n📋 Todas las celdas de la fila 3:")
    for col in range(1, 15):
        val = ws.cell(row=3, column=col).value
        col_letter = openpyxl.utils.get_column_letter(col)
        if val:
            print(f"   {col_letter}3: '{val}'")
    
    # Verificar E4
    e4_value = ws.cell(row=4, column=5).value
    print(f"\n🔍 Celda E4 (fila 4, columna 5):")
    print(f"   Valor: '{e4_value}'")
    
    # Verificar I3
    i3_value = ws.cell(row=3, column=9).value
    print(f"\n🔍 Celda I3 (fila 3, columna 9):")
    print(f"   Valor: '{i3_value}'")
    
    wb.close()
    print("\n" + "=" * 70)
else:
    print(f"❌ Archivo no encontrado: {archivo}")

# Ahora verificar también el archivo real generado
print("\n" + "=" * 70)
print("VERIFICANDO ARCHIVO REAL DE EXPORTACIÓN")
print("=" * 70)

archivo_real = Path("workers/fer/012026/result")
if archivo_real.exists():
    # Buscar el último archivo exportado
    archivos = sorted(archivo_real.glob("export_*.xlsx"))
    if archivos:
        ultimo = archivos[-1]
        print(f"\nÚltimo archivo: {ultimo.name}")
        
        wb = openpyxl.load_workbook(ultimo)
        ws = wb.active
        
        f3_value = ws.cell(row=3, column=6).value
        e4_value = ws.cell(row=4, column=5).value
        i3_value = ws.cell(row=3, column=9).value
        
        print(f"   F3 (trabajador): '{f3_value}'")
        print(f"   E4 (rango fechas): '{e4_value}'")
        print(f"   I3 (payment type): '{i3_value}'")
        
        wb.close()
    else:
        print("No se encontraron archivos export_*.xlsx")
else:
    print("Directorio no encontrado")

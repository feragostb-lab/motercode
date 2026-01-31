"""Comparar los archivos exportados para verificar que tienen la misma estructura."""
import openpyxl
from pathlib import Path

def analizar_archivo(filepath):
    """Analiza un archivo Excel y retorna información sobre su estructura."""
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active
    
    info = {
        'max_row': ws.max_row,
        'max_col': ws.max_column,
        'formula_suma': None,
        'ultima_fila_datos': None,
        'num_reglas_condicionales': 0,
        'primer_dato_col3': None,
        'ultimo_dato_col3': None
    }
    
    # Detectar última fila con datos en columna 3
    for row in range(ws.max_row, 0, -1):
        val = ws.cell(row=row, column=3).value
        if val is not None and str(val).strip() != "":
            info['ultima_fila_datos'] = row
            info['ultimo_dato_col3'] = val
            break
    
    # Detectar primera fila de datos (después de cabecera)
    for row in range(1, ws.max_row + 1):
        val = ws.cell(row=row, column=3).value
        if val is not None and str(val).strip() != "" and row > 1:
            info['primer_dato_col3'] = val
            break
    
    # Buscar fórmula de suma
    if info['ultima_fila_datos']:
        fila_siguiente = info['ultima_fila_datos'] + 1
        if fila_siguiente <= ws.max_row:
            formula = ws.cell(row=fila_siguiente, column=9).value
            if formula and str(formula).startswith('=SUM'):
                info['formula_suma'] = formula
    
    # Contar reglas de formato condicional
    if hasattr(ws.conditional_formatting, '_cf_rules'):
        info['num_reglas_condicionales'] = len(ws.conditional_formatting._cf_rules)
    
    wb.close()
    return info

def main():
    archivos = {
        'excel_test.py': Path('doc/examples/resultado_copia.xlsx'),
        'export_service': Path('doc/examples/test_export_resultado.xlsx')
    }
    
    print("=" * 70)
    print("COMPARACIÓN DE ARCHIVOS EXPORTADOS")
    print("=" * 70)
    
    resultados = {}
    
    for nombre, path in archivos.items():
        if not path.exists():
            print(f"\n❌ {nombre}: Archivo no encontrado - {path}")
            continue
        
        print(f"\n📄 Analizando: {nombre}")
        print(f"   Ruta: {path}")
        
        info = analizar_archivo(path)
        resultados[nombre] = info
        
        print(f"   ├─ Filas totales: {info['max_row']}")
        print(f"   ├─ Columnas totales: {info['max_col']}")
        print(f"   ├─ Primera fila de datos (col 3): fila con valor '{info['primer_dato_col3']}'")
        print(f"   ├─ Última fila con datos: {info['ultima_fila_datos']}")
        print(f"   ├─ Último valor (col 3): {info['ultimo_dato_col3']}")
        print(f"   ├─ Fórmula de suma: {info['formula_suma']}")
        print(f"   └─ Reglas condicionales: {info['num_reglas_condicionales']}")
    
    print("\n" + "=" * 70)
    print("RESUMEN DE COMPARACIÓN")
    print("=" * 70)
    
    if len(resultados) == 2:
        keys = list(resultados.keys())
        r1 = resultados[keys[0]]
        r2 = resultados[keys[1]]
        
        comparaciones = [
            ('Filas totales', r1['max_row'], r2['max_row']),
            ('Última fila datos', r1['ultima_fila_datos'], r2['ultima_fila_datos']),
            ('Reglas condicionales', r1['num_reglas_condicionales'], r2['num_reglas_condicionales']),
            ('Fórmula suma presente', bool(r1['formula_suma']), bool(r2['formula_suma']))
        ]
        
        print()
        for metrica, val1, val2 in comparaciones:
            match = "✅" if val1 == val2 else "❌"
            print(f"{match} {metrica}: {keys[0]}={val1}, {keys[1]}={val2}")
        
        print("\n" + "=" * 70)
        if all(val1 == val2 for _, val1, val2 in comparaciones):
            print("✅ ¡ÉXITO! Ambos archivos tienen la misma estructura")
        else:
            print("⚠️  Los archivos tienen diferencias en su estructura")
        print("=" * 70)

if __name__ == "__main__":
    main()

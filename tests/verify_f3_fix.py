"""Verify that F3 cell contains worker name after export."""
import openpyxl
from pathlib import Path

# Check the most recent export file
export_file = Path("exports/all_transactions_complete_20260130_133401.xlsx")

if not export_file.exists():
    print(f"❌ File not found: {export_file}")
    print("Please create a new export first.")
else:
    print(f"📊 Checking file: {export_file}")
    
    wb = openpyxl.load_workbook(export_file)
    ws = wb.active
    
    # Check F3 cell
    f3_value = ws.cell(row=3, column=6).value  # Column F = 6
    print(f"\n📍 Cell F3 value: {f3_value}")
    
    if f3_value:
        print(f"✅ F3 contains: '{f3_value}'")
    else:
        print("❌ F3 is empty or None")
    
    # Also check E4 and I3 for completeness
    e4_value = ws.cell(row=4, column=5).value  # Column E = 5
    i3_value = ws.cell(row=3, column=9).value  # Column I = 9
    
    print(f"\n📍 Cell E4 (Date range): {e4_value}")
    print(f"📍 Cell I3 (Payment type): {i3_value}")
    
    wb.close()
    
    print("\n" + "="*50)
    print("To test the fix:")
    print("1. Ensure you have an active worker and period")
    print("2. Export from the dashboard: 'Export All Transactions Complete'")
    print("3. Check the F3 cell in the new file")
    print("="*50)

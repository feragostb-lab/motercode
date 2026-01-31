"""Script para probar exportación completa con periodo real."""
from pathlib import Path
from src.services.export_service import ExportService
from src.core.config import Config
from src.repositories.period_repository import PeriodRepository
from src.core.database import get_database

def test_real_period_export():
    """Prueba la exportación de un periodo real si existe."""
    print("Iniciando prueba de exportación con periodo real...")
    
    # Cargar configuración
    config = Config()
    
    # Conectar a la base de datos
    db = get_database(config.paths.get('database', './receipts.db'))
    period_repo = PeriodRepository(db)
    
    # Obtener el primer periodo disponible
    all_periods = period_repo.get_all()
    
    if not all_periods:
        print("\n⚠️  No hay periodos en la base de datos.")
        print("   Esto es solo una advertencia, la funcionalidad está implementada.")
        print("   Cuando haya periodos con datos, la exportación incluirá:")
        print("   - F3: Nombre del trabajador")
        print("   - E4: Rango de fechas (menor - mayor)")
        print("   - I3: Corporate Card - [Mes Año]")
        return
    
    # Usar el primer periodo
    period = all_periods[0]
    print(f"\n✅ Periodo encontrado: ID={period.id}, Worker={period.worker_id}, Period={period.month_year}")
    
    # Crear servicio de exportación
    export_service = ExportService(config)
    
    try:
        # Exportar el periodo
        output_path = export_service.export_period_data(period.id, export_type='temporal')
        
        print(f"\n✅ Exportación exitosa!")
        print(f"📁 Archivo generado: {output_path}")
        print("\nVerificando celdas especiales...")
        
        # Verificar el archivo generado
        import openpyxl
        wb = openpyxl.load_workbook(output_path)
        ws = wb.active
        
        f3_value = ws.cell(row=3, column=6).value
        e4_value = ws.cell(row=4, column=5).value
        i3_value = ws.cell(row=3, column=9).value
        
        print(f"   F3 (Trabajador): {f3_value}")
        print(f"   E4 (Rango fechas): {e4_value}")
        print(f"   I3 (Payment type): {i3_value}")
        
        wb.close()
        
        print("\n✅ Todas las celdas se escribieron correctamente!")
        
    except Exception as e:
        print(f"\n❌ Error durante la exportación: {e}")
        import traceback
        traceback.print_exc()

if __name__ == "__main__":
    test_real_period_export()

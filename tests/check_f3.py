"""Script para verificar el contenido de F3 en cualquier archivo."""
import openpyxl
from pathlib import Path
import sys

def verificar_archivo(filepath):
    """Verifica el contenido de F3 en un archivo."""
    if not filepath.exists():
        print(f"❌ Archivo no encontrado: {filepath}")
        return
    
    print(f"\n{'='*70}")
    print(f"VERIFICANDO: {filepath.name}")
    print(f"{'='*70}")
    
    wb = openpyxl.load_workbook(filepath)
    ws = wb.active
    
    # Verificar F3
    f3_value = ws.cell(row=3, column=6).value
    print(f"\n✨ Celda F3 (Nombre trabajador):")
    print(f"   Contenido: '{f3_value}'")
    print(f"   Tipo: {type(f3_value).__name__}")
    
    if f3_value and f3_value != '[worker name]':
        print(f"   Estado: ✅ TIENE NOMBRE DEL TRABAJADOR")
    elif f3_value == '[worker name]':
        print(f"   Estado: ⚠️  Tiene el placeholder original de la plantilla")
    else:
        print(f"   Estado: ❌ VACÍA O NULA")
    
    # Mostrar también E4 e I3 para contexto
    e4_value = ws.cell(row=4, column=5).value
    i3_value = ws.cell(row=3, column=9).value
    
    print(f"\n📊 Otras celdas:")
    print(f"   E4 (Rango fechas): '{e4_value}'")
    print(f"   I3 (Payment type): '{i3_value}'")
    
    # Mostrar datos
    primera_fila_datos = ws.cell(row=8, column=3).value
    if primera_fila_datos:
        print(f"\n📋 Primera fila de datos (fila 8):")
        print(f"   ID: {ws.cell(row=8, column=3).value}")
        print(f"   Fecha: {ws.cell(row=8, column=4).value}")
        print(f"   Amount: {ws.cell(row=8, column=9).value}")
    
    wb.close()
    print(f"{'='*70}\n")

if __name__ == "__main__":
    print("\n🔍 VERIFICADOR DE ARCHIVOS EXCEL - CELDA F3")
    print("="*70)
    
    # Si se pasa un argumento, verificar ese archivo
    if len(sys.argv) > 1:
        archivo = Path(sys.argv[1])
        verificar_archivo(archivo)
    else:
        # Verificar archivos comunes
        archivos_a_verificar = [
            Path("doc/examples/test_export_resultado.xlsx"),
            Path("doc/examples/resultado_copia.xlsx"),
        ]
        
        # Buscar también archivos en workers/
        workers_dir = Path("workers")
        if workers_dir.exists():
            for worker_folder in workers_dir.iterdir():
                if worker_folder.is_dir():
                    for period_folder in worker_folder.iterdir():
                        if period_folder.is_dir():
                            result_dir = period_folder / "result"
                            if result_dir.exists():
                                for excel_file in sorted(result_dir.glob("*.xlsx")):
                                    archivos_a_verificar.append(excel_file)
        
        print(f"Verificando {len(archivos_a_verificar)} archivo(s)...\n")
        
        for archivo in archivos_a_verificar:
            if archivo.exists():
                verificar_archivo(archivo)
    
    print("\n💡 TIP: Si F3 aparece vacía en Excel pero este script muestra el nombre,")
    print("   cierra completamente Excel y vuelve a abrir el archivo.")
